#!/usr/bin/env python3
"""Generate the five-minute assignment presentation from validated artifacts."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(24, 53, 82)
TEAL = RGBColor(8, 163, 184)
GOLD = RGBColor(245, 166, 35)
GREEN = RGBColor(5, 150, 105)
LIGHT = RGBColor(242, 246, 249)
MUTED = RGBColor(86, 105, 124)
WHITE = RGBColor(255, 255, 255)
FONT = "Noto Sans CJK KR"


def add_text(slide, text, left, top, width, height, size=24, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.25, 12.0, 0.65, size=28, bold=True)
    slide.shapes.add_shape(1, Inches(0.65), Inches(0.95), Inches(1.05), Inches(0.06)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = GOLD
    slide.shapes[-1].line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 1.9, 0.78, 10.5, 0.35, size=12, color=MUTED)


def add_footer(slide, number):
    add_text(slide, "Robotics Simulation · Week 03", 0.65, 7.08, 5.0, 0.24, size=9, color=MUTED)
    add_text(slide, str(number), 12.2, 7.05, 0.45, 0.28, size=10, bold=True, color=MUTED, align=PP_ALIGN.RIGHT)


def add_bullet_list(slide, bullets, left, top, width, height, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = NAVY
        paragraph.space_after = Pt(10)
        paragraph.text = f"•  {bullet}"
    return box


def add_table(slide, rows, left, top, width, height, col_widths=None):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)).table
    if col_widths:
        for index, value in enumerate(col_widths):
            table.columns[index].width = Inches(value)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if row_index == 0 else (LIGHT if row_index % 2 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = FONT
                paragraph.font.size = Pt(13 if row_index else 14)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = WHITE if row_index == 0 else NAVY
                paragraph.alignment = PP_ALIGN.CENTER
    return table


def load_summary(path: Path):
    with path.open(encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def metric(summary, variant, scenario):
    for row in summary:
        if row["variant"] == variant and row["scenario"] == scenario:
            return f"{float(row['return_mean']):.1f} ± {float(row['return_std']):.1f}"
    return "pending"


def mean_value(summary, variant, scenario):
    for row in summary:
        if row["variant"] == variant and row["scenario"] == scenario:
            return float(row["return_mean"])
    raise ValueError(f"missing result for {variant}/{scenario}")


def percent_change(new, reference):
    return 100.0 * (new - reference) / abs(reference)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--github-url", default="PUBLIC_GITHUB_URL_PENDING")
    parser.add_argument("--output", type=Path, default=ROOT / "report" / "week03_ant_robust_report.pptx")
    args = parser.parse_args()

    summary_path = ROOT / "artifacts" / "evaluations" / "evaluation_summary.csv"
    summary = load_summary(summary_path)
    scenarios = ("low_friction", "heavy", "push")
    baseline_id = mean_value(summary, "baseline", "id")
    robust_id = mean_value(summary, "robust", "id")
    baseline_ood = sum(mean_value(summary, "baseline", item) for item in scenarios) / len(scenarios)
    robust_ood = sum(mean_value(summary, "robust", item) for item in scenarios) / len(scenarios)
    improved_ood = sum(
        mean_value(summary, "robust", item) > mean_value(summary, "baseline", item)
        for item in scenarios
    )
    hypothesis = "지지" if improved_ood >= 2 and robust_id >= 0.9 * baseline_id else "기각"
    learning_plot = ROOT / "artifacts" / "plots" / "training_curves.png"
    evaluation_plot = ROOT / "artifacts" / "plots" / "evaluation_returns.png"
    video_thumbnail = ROOT / "artifacts" / "videos" / "thumbnails" / "comparison_low_friction_seed42_frame480.png"
    for required in (learning_plot, evaluation_plot, video_thumbnail):
        if not required.is_file():
            raise FileNotFoundError(required)

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    add_text(slide, "처음 보는 환경에서도\n잘 걷는 Ant 만들기", 0.8, 1.2, 8.6, 1.8, size=34, bold=True, color=WHITE)
    add_text(slide, "PPO domain randomization · 100-env evaluation", 0.85, 3.2, 8.5, 0.55, size=19, color=RGBColor(185, 230, 235))
    add_text(slide, "Robotics Simulation Week 03", 0.85, 5.8, 7.0, 0.4, size=15, color=WHITE)
    add_text(slide, "TEAM MEMBERS: TBD", 0.85, 6.25, 7.0, 0.4, size=13, color=RGBColor(190, 200, 212))
    add_text(slide, "01", 11.6, 5.2, 1.0, 1.0, size=42, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

    slide = deck.slides.add_slide(blank)
    add_title(slide, "질문과 공정한 실험 설계", "한 번의 최고 점수가 아니라 같은 budget의 일반화 비교")
    add_bullet_list(
        slide,
        [
            "가설: 제한적 물리 랜덤화가 ID 성능을 크게 잃지 않고 OOD 생존성과 return을 높인다.",
            "통제: 60D observation, 8D action, PPO 구조, 4096 envs, 32 steps/env, 1000 iterations.",
            "반복: baseline/robust 각 3 seeds (42·43·44), evaluation은 seed 24의 100 envs.",
            "판정: low-friction·payload·push 중 둘 이상에서 개선되고 ID collapse가 없어야 한다.",
        ],
        0.9,
        1.45,
        11.6,
        4.1,
        size=19,
    )
    add_table(
        slide,
        [["Budget", "Value"], ["Transitions/run", "1000 × 4096 × 32 = 131,072,000"], ["Checkpoint", "model_999.pt"], ["Metric", "first-episode return mean ± population std"]],
        1.25,
        5.25,
        10.8,
        1.35,
        [3.0, 7.8],
    )
    add_footer(slide, 2)

    slide = deck.slides.add_slide(blank)
    add_title(slide, "무엇을 바꿨는가", "interface는 그대로, training distribution만 확장")
    rows = [
        ["Variant", "Training distribution", "Purpose"],
        ["Baseline", "course Isaac-Ant-v0", "reference"],
        ["Friction", "μs 0.45–1.35, μd 0.35–1.15", "single-factor ablation"],
        ["Robust", "friction + mass/COM + reset + noise + push", "combined generalization"],
    ]
    add_table(slide, rows, 0.75, 1.35, 11.85, 2.45, [2.0, 6.1, 3.75])
    add_text(slide, "Public holdouts", 0.8, 4.1, 4.0, 0.45, size=21, bold=True, color=TEAL)
    add_bullet_list(
        slide,
        ["Low friction: 0.30 / 0.25 (training range 밖)", "Heavy: torso ×1.30 + off-center COM", "Push: ±0.80 m/s impulse every 3–5 s"],
        0.9,
        4.6,
        7.3,
        1.7,
        size=18,
    )
    add_text(slide, "Hidden test는 발표 당일 공개\n→ public holdout은 대체 점수가 아니라 반증 도구", 8.5, 4.7, 3.7, 1.25, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_footer(slide, 3)

    slide = deck.slides.add_slide(blank)
    add_title(slide, "학습은 안정적으로 수렴했는가", "TensorBoard Train/mean_reward · 25-iteration moving average")
    slide.shapes.add_picture(str(learning_plot), Inches(1.15), Inches(1.25), width=Inches(11.0))
    add_footer(slide, 4)

    slide = deck.slides.add_slide(blank)
    add_title(slide, "100개 환경 정량 평가", "각 vectorized env의 첫 episode만 누적")
    slide.shapes.add_picture(str(evaluation_plot), Inches(0.65), Inches(1.25), width=Inches(8.0))
    rows = [["Scenario", "Baseline", "Robust"]]
    for scenario, label in (("id", "ID"), ("low_friction", "Low μ"), ("heavy", "Heavy"), ("push", "Push")):
        rows.append([label, metric(summary, "baseline", scenario), metric(summary, "robust", scenario)])
    add_table(slide, rows, 8.75, 1.55, 4.0, 3.4, [1.1, 1.45, 1.45])
    add_text(slide, "표의 ±는 100×seed pooled episode population std", 8.85, 5.15, 3.8, 0.6, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    slide.shapes.add_picture(str(video_thumbnail), Inches(8.95), Inches(5.72), width=Inches(3.6))
    add_footer(slide, 5)

    slide = deck.slides.add_slide(blank)
    add_title(slide, "결론 · 재현성 · 한계", "높은 점수 하나보다 설계와 반증 가능성을 남긴다")
    add_bullet_list(
        slide,
        [
            (
                f"가설 {hypothesis}: Robust는 ID {percent_change(robust_id, baseline_id):+.1f}%, "
                f"공개 OOD 평균 {percent_change(robust_ood, baseline_ood):+.1f}% (baseline 대비)."
            ),
            "README에 exact command, config, log, checkpoint SHA-256, 100-env JSON을 공개한다.",
            "영상: baseline과 robust를 같은 seed/holdout에서 비교한다.",
            "한계: hidden test 미사용, 3 seeds, terrain/latency shift는 미포함.",
        ],
        0.85,
        1.45,
        11.7,
        3.3,
        size=19,
    )
    add_text(slide, "Evaluation command", 0.9, 5.05, 3.0, 0.35, size=16, bold=True, color=TEAL)
    add_text(slide, "./scripts/run_evaluate.sh --num_envs 100 --seed 24 --checkpoint …", 0.9, 5.45, 11.4, 0.45, size=14, color=NAVY)
    add_text(slide, args.github_url, 0.9, 6.25, 11.4, 0.4, size=15, bold=True, color=GREEN)
    add_footer(slide, 6)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
